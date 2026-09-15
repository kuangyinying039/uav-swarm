#!/usr/bin/env python3
"""Replace EQ* placeholders with native Office Math (OMML / MathType-compatible)."""
from __future__ import annotations

import sys
from lxml import etree
from pptx import Presentation
M = "http://schemas.openxmlformats.org/officeDocument/2006/math"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
PNS = "http://schemas.openxmlformats.org/presentationml/2006/main"


etree.register_namespace("m", M)


def me(tag):
    return f"{{{M}}}{tag}"


def ae(tag):
    return f"{{{A}}}{tag}"


def r(text, sty="i"):
    node = etree.Element(me("r"))
    rpr = etree.SubElement(node, me("rPr"))
    sty_el = etree.SubElement(rpr, me("sty"))
    sty_el.set(f"{{{M}}}val", sty)
    t = etree.SubElement(node, me("t"))
    if text.startswith(" ") or text.endswith(" "):
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    return node


def kids(tag, items):
    node = etree.Element(me(tag))
    seq = items if isinstance(items, (list, tuple)) else [items]
    for c in seq:
        node.append(c)
    return node


def sub(base, subscript):
    node = etree.Element(me("sSub"))
    node.append(kids("e", base if isinstance(base, (list, tuple)) else [base]))
    node.append(kids("sub", subscript if isinstance(subscript, (list, tuple)) else [subscript]))
    return node


def sup(base, superscript):
    node = etree.Element(me("sSup"))
    node.append(kids("e", base if isinstance(base, (list, tuple)) else [base]))
    node.append(kids("sup", superscript if isinstance(superscript, (list, tuple)) else [superscript]))
    return node


def frac(num, den):
    node = etree.Element(me("f"))
    etree.SubElement(node, me("fPr"))
    node.append(kids("num", num if isinstance(num, (list, tuple)) else [num]))
    node.append(kids("den", den if isinstance(den, (list, tuple)) else [den]))
    return node


def dmath(*inner):
    node = etree.Element(me("d"))
    dpr = etree.SubElement(node, me("dPr"))
    beg = etree.SubElement(dpr, me("begChr"))
    beg.set(f"{{{M}}}val", "(")
    end = etree.SubElement(dpr, me("endChr"))
    end.set(f"{{{M}}}val", ")")
    node.append(kids("e", inner))
    return node


def func(name, *args):
    node = etree.Element(me("func"))
    etree.SubElement(node, me("funcPr"))
    fname = etree.SubElement(node, me("fName"))
    fname.append(kids("e", [r(name, "p")]))
    node.append(kids("e", args))
    return node


def om(children):
    math = etree.Element(me("oMath"))
    for c in children:
        math.append(c)
    para = etree.Element(me("oMathPara"))
    pr = etree.SubElement(para, me("oMathParaPr"))
    jc = etree.SubElement(pr, me("jc"))
    jc.set(f"{{{M}}}val", "center")
    para.append(math)
    return para


def eq_motion():
    return om([
        sub(r("p"), [r("i"), r(",", "p"), r("t"), r("+", "p"), r("1", "p")]),
        r(" = ", "p"),
        sub(r("Π", "p"), r("Ω")),
        dmath(
            sub(r("p"), [r("i"), r(",", "p"), r("t")]),
            r(" + ", "p"),
            r("v", "i"),
            sub(r("Δt", "i"), r("t")),
            r(" ", "p"),
            r("d"),
            dmath(sub(r("a"), [r("i"), r(",", "p"), r("t")])),
        ),
    ])


def eq_link():
    return om([
        sub(r("A"), [r("i"), r(",", "p"), r("j"), r(",", "p"), r("t")]),
        r(" = ", "p"),
        r("1", "p"),
        dmath(
            r("∥", "p"),
            sub(r("p"), [r("i"), r(",", "p"), r("t")]),
            r(" − ", "p"),
            sub(r("p"), [r("j"), r(",", "p"), r("t")]),
            r("∥", "p"),
            r(" ≤ ", "p"),
            sub(r("R"), r("c")),
        ),
        r(" ", "p"),
        sub(r("z"), [r("i"), r(",", "p"), r("t")]),
        sub(r("z"), [r("j"), r(",", "p"), r("t")]),
        sub(r("u"), [r("i"), r(",", "p"), r("t")]),
        sub(r("u"), [r("j"), r(",", "p"), r("t")]),
    ])


def eq_p():
    return om([
        sub(r("P"), [r("i"), r(",", "p"), r("t")]),
        r("(x)", "p"),
        r(" = ", "p"),
        frac(
            r("1", "p"),
            [
                r("1", "p"),
                r(" + ", "p"),
                func("exp", r("−", "p"), sub(r("L"), [r("i"), r(",", "p"), r("t")]), r("(x)", "p")),
            ],
        ),
    ])


def eq_w():
    return om([
        sub(r("w"), [r("i"), r(",", "p"), r("j"), r(",", "p"), r("t")]),
        r(" = ", "p"),
        frac(
            [
                sup(
                    r("e"),
                    [
                        r("−", "p"),
                        sub(r("β"), r("d")),
                        r("∥", "p"),
                        sub(r("p"), r("i")),
                        r(" − ", "p"),
                        sub(r("p"), r("j")),
                        r("∥", "p"),
                    ],
                ),
            ],
            [r("1", "p"), r(" + ", "p"), sub(r("g"), [r("j"), r(",", "p"), r("t")])],
        ),
    ])


def eq_jstar():
    return om([
        sup(r("j"), r("⋆", "p")),
        r(" = ", "p"),
        r("arg", "p"),
        r(" max", "p"),
        sub(r(""), r("j")),
        r(" ", "p"),
        sub(r("w"), [r("i"), r(",", "p"), r("j"), r(",", "p"), r("t")]),
        r(" |", "p"),
        sub(r("L"), [r("j"), r(",", "p"), r("t")]),
        r("(x)|", "p"),
    ])


def eq_att():
    inner = [
        frac(
            [sub(r("Q"), r("i")), sup(sub(r("K"), r("k")), r("⊤", "p"))],
            [r("√", "p"), r("d")],
        ),
        r(" + ", "p"),
        func("log", sub(r("χ"), [r("r"), r(",", "p"), r("i"), r(",", "p"), r("k")])),
        r(" − ", "p"),
        sub(r("κ"), r("r")),
        sub(r("d"), [r("i"), r("k")]),
    ]
    return om([
        sub(r("α"), [r("r"), r(",", "p"), r("i"), r(",", "p"), r("k")]),
        r(" = ", "p"),
        func("softmax", dmath(*inner)),
    ])


def eq_r():
    return om([
        sub(r("r"), r("t")),
        r(" = ", "p"),
        sub(r("λ"), r("D")),
        sub(r("ΔD"), r("t")),
        r(" + ", "p"),
        sub(r("λ"), r("C")),
        sub(r("ΔC"), r("t")),
        r(" − ", "p"),
        sub(r("λ"), [r("safe", "p")]),
        sub(r("c"), r("t")),
    ])


EQS = {
    "EQMOTION": eq_motion,
    "EQLINK": eq_link,
    "EQP": eq_p,
    "EQW": eq_w,
    "EQJSTAR": eq_jstar,
    "EQATT": eq_att,
    "EQREWARD": eq_r,
}


def replace_placeholder(sld, key, builder):
    found = 0
    for t in sld.iter(ae("t")):
        if t.text and t.text.strip() == key:
            p = t
            while p is not None and p.tag != ae("p"):
                p = p.getparent()
            if p is None:
                continue
            pPr = p.find(ae("pPr"))
            for child in list(p):
                p.remove(child)
            if pPr is not None:
                p.insert(0, pPr)
            p.append(builder())
            found += 1
    return found


def ensure_math_ns(sld):
    sld.set("{http://www.w3.org/2000/xmlns/}m", M)


def main(path):
    prs = Presentation(path)
    total = 0
    for slide in prs.slides:
        for key, builder in EQS.items():
            total += replace_placeholder(slide._element, key, builder)
    prs.save(path)
    print(f"injected {total} OMML equations into {path}")
    if total < 7:
        print("WARNING: expected at least 7 equation placeholders", file=sys.stderr)
        sys.exit(1)


if __name__ == "__main__":
    main(sys.argv[1])
